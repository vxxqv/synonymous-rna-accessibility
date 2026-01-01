"""Authenticate metadata for the proposed K562 comparison; no assay numeric reads."""
from pathlib import Path
import csv, hashlib, json, re, zipfile
from openpyxl import load_workbook

if not __debug__:
    raise RuntimeError("Optimized execution refused")
BASE = Path(__file__).resolve().parent
OLD = BASE.parents[1] / "research_cycle8/cprime_gate"
WORKBOOK = OLD / "supp_Supplementary Data 6.xlsx"
SHEET = "validation K562"
LABELS = {
    "edit_only": "R1207R_(A/G)TTG",
    "marker_only": "R1207R_ATT(G/A)",
    "combined_A": "R1207R_(A/G)TT(G/A)",
}
FAMILIES = {
    "edit_only": "K562_wo_marker ",
    "marker_only": "K562_wo_edit ",
    "combined_A": "K562_marker ",
}
with zipfile.ZipFile(OLD / "41467_2026_78344_MOESM3_ESM.zip") as archive:
    entries = [n for n in archive.namelist() if n.endswith("Supplementary Data 6.xlsx")]
    if len(entries) != 1 or archive.read(entries[0]) != WORKBOOK.read_bytes():
        raise RuntimeError("Source Data6 extraction does not match official cached container")
    member = entries[0]

book = load_workbook(WORKBOOK, read_only=True, data_only=False)
sheet = book[SHEET]
header = {}
# These exact first-row cells are schema, not observations.
for col in [3, 6, 7, 8, 9, 10, 11]:
    cell = sheet.cell(1, col)
    if not isinstance(cell.value, str) or cell.data_type == "f":
        raise RuntimeError("Unexpected header schema")
    header[cell.column_letter] = cell.value
expected = {"C": "Paper_name", "F": "K562_gDNA_1_RPM", "G": "K562_gDNA_2_RPM",
            "H": "K562_gDNA_3_RPM", "I": "K562_cDNA_1_RPM", "J": "K562_cDNA_2_RPM",
            "K": "K562_cDNA_3_RPM"}
if header != expected:
    raise RuntimeError("Unexpected Data6 schema")
row_map = {}
all_labels = []
for cells in sheet.iter_rows(min_row=2, min_col=3, max_col=3):
    cell = cells[0]
    value = cell.value
    if value is None:
        continue
    if not isinstance(value, str) or cell.data_type == "f":
        raise RuntimeError("Nonliteral label in identity column")
    all_labels.append({"row": cell.row, "label": value})
    for state, label in LABELS.items():
        if value == label:
            if state in row_map:
                raise RuntimeError("Duplicate selected identity")
            row_map[state] = cell.row
book.close()
if set(row_map) != set(LABELS):
    raise RuntimeError("Missing required identity")

with (OLD / "ena_samples.tsv").open(encoding="utf-8-sig", newline="") as stream:
    runs = list(csv.DictReader(stream, delimiter="\t"))
samples = {}
excluded_earlier = []
for state, label in LABELS.items():
    matched = []
    for r in runs:
        alias = r["experiment_alias"]
        if label not in alias:
            continue
        if alias.startswith(FAMILIES[state]):
            m = re.search(r"_(\d)$", alias)
            mode = "cDNA" if "cDNA" in alias else "gDNA" if "gDNA" in alias else None
            if not m or mode is None:
                raise RuntimeError("Ambiguous modality/replicate label")
            matched.append({**r, "metadata_modality": mode, "source_replicate_label": int(m.group(1))})
        elif alias.startswith("K562_"):
            excluded_earlier.append({"state": state, "reason": "Earlier unprefixed technical family; never pool",
                                     **r})
    shape = sorted((r["metadata_modality"], r["source_replicate_label"]) for r in matched)
    if shape != sorted((m, i) for m in ["cDNA", "gDNA"] for i in [1, 2, 3]):
        raise RuntimeError(f"Incomplete or duplicate source metadata family: {state}")
    samples[state] = sorted(matched, key=lambda r: (r["metadata_modality"], r["source_replicate_label"]))

validation = json.loads((OLD / "published_validation_sequence_metadata.json").read_text())
library = json.loads((OLD / "ex27_selected_metadata.json").read_text())["library"]
id_by_state = {
    "edit_only": "Synonymous_edit_29560144_wo_marker",
    "marker_only": "Synonymous_edit_29560144_wo_intended_edit",
    "combined_A": "Synonymous_edit_29560144",
}
state_metadata = {}
for state, variant in id_by_state.items():
    if state == "combined_A":
        match = [r for r in library if r["Sample ID"] == variant]
        if len(match) != 1:
            raise RuntimeError("Combined state not unique")
        dna, rna = match[0]["Search sequence gDNA"], match[0]["Search sequence cDNA"]
        wt = match[0]["WT sequence gDNA"]
    else:
        match = [r for r in validation if r["id"] == variant]
        if len(match) != 1:
            raise RuntimeError("Single state not unique")
        dna, rna, wt = match[0]["search_DNA"], match[0]["search_RNA"], match[0]["WT_DNA"]
    if dna != rna or len(dna) != len(wt):
        raise RuntimeError("Local sequence mismatch")
    state_metadata[state] = {
        "source_id": variant, "paper_label": LABELS[state], "sheet": SHEET,
        "row": row_map[state], "label_cell": f"C{row_map[state]}",
        "gDNA_RPM_cells": [f"{c}{row_map[state]}" for c in ["F", "G", "H"]],
        "cDNA_RPM_cells": [f"{c}{row_map[state]}" for c in ["I", "J", "K"]],
        "local_DNA_search_word": dna, "local_cDNA_search_word": rna,
        "local_WT_reference_word": wt,
        "local_1based_differences": [
            {"position": i, "reference": a, "state": b}
            for i, (a, b) in enumerate(zip(wt, dna), 1) if a != b
        ],
        "source_run_metadata": samples[state],
    }

contract = {
    "status": "QUALIFIED_DESCRIPTIVE_CONTRACT_READY_NUMERIC_ACCESS_CLOSED",
    "workbook_path": str(WORKBOOK),
    "workbook_sha256": hashlib.sha256(WORKBOOK.read_bytes()).hexdigest(),
    "official_container_url": "https://pmc-oa-opendata.s3.amazonaws.com/PMC13675207.1/41467_2026_78344_MOESM3_ESM.zip",
    "official_container_sha256": hashlib.sha256((OLD / "41467_2026_78344_MOESM3_ESM.zip").read_bytes()).hexdigest(),
    "official_member": member, "sheet": SHEET, "header_schema": header,
    "state_metadata": state_metadata,
    "all_sheet_identity_labels": all_labels,
    "included_named_source_runs": sum(len(x) for x in samples.values()),
    "excluded_earlier_technical_run_metadata": excluded_earlier,
    "excluded_states": ["marker_GtoT: unresolved count-word mismatch",
                        "marker_GtoC: outside the fixed three-state question",
                        "WT: local reference word exists, compatible measured WT RNA response not recovered"],
    "source_cell_model": "K562_PE_MLH1 as described for arrayed experiments",
    "source_transcript_annotation": "ENST00000358273; intended c.3621A>G and marker c.3624G>A",
    "biological_independence_authenticated": False,
    "cross_arm_culture_pairing_authenticated": False,
    "arm_specific_harvest_and_clone_mapping_authenticated": False,
    "Data6_to_ENA_linkage": "Exact condition/modality/replicate-name compatibility; no deposited explicit cell-to-run provenance key recovered",
    "raw_counts_RPM_scores_or_missingness_inspected": False,
    "no_new_outcomes_interpreted": True,
}
target = BASE / "three_state_contract.json"
if target.exists():
    raise FileExistsError(target)
target.write_text(json.dumps(contract, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"status": contract["status"], "workbook_sha256": contract["workbook_sha256"],
                  "rows": row_map, "included_named_source_runs": contract["included_named_source_runs"],
                  "excluded_earlier_technical_runs": len(excluded_earlier),
                  "local_differences": {s:r["local_1based_differences"] for s,r in state_metadata.items()}}, indent=2))
